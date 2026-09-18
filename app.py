"""
Cruscotto Grandi Apparecchiature — AUSL Toscana Nord Ovest.
Avvio:  streamlit run app.py
Sorgente dati di default: estrazione NSIS (offline). Commutabile su API AT2.0
dalla sidebar quando l'ambiente Metis sarà attivo.
"""
from io import BytesIO

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import data_source as ds

COLORE_HEADER = "366092"       # blu — stesso stile del report Direzione della dashboard PNRR
COLORI_SEMAFORO_XLSX = {
    "Vetusto": "F8D7DA", "In avvicinamento": "FFF3CD", "OK": "D1E7DD",
}
ETICHETTE_SORGENTE = {
    "nsis": "estrazione NSIS (offline)",
    "cache": "API AT2.0 — cache locale (parco.parquet)",
    "api": "API AT2.0 — live",
}
def _euro(v) -> str:
    return "n.d." if pd.isna(v) else f"€ {v:,.0f}"


COLONNE_ELENCO_EXPORT = [
    "n_inventario", "tipologia", "componente", "fabbricante", "modello", "zona", "sede",
    "anno_collaudo", "anzianita_anni", "stato_vetusta", "modalita_acquisizione",
    "valore_economico", "scadenza_contratto", "scadenza_vse", "scadenza_cq",
]

st.set_page_config(page_title="Grandi Apparecchiature — AUSL TNO",
                   page_icon="🩻", layout="wide")


# --------------------------------------------------------------------------- #
# Caricamento dati (con cache)
# --------------------------------------------------------------------------- #
@st.cache_data(show_spinner="Caricamento parco apparecchiature…")
def carica(sorgente: str, nsis_path: str, max_records: int | None = None) -> pd.DataFrame:
    return ds.carica(sorgente, nsis_path=nsis_path, max_records=max_records)


# --------------------------------------------------------------------------- #
# Helper: report Direzione in Excel (intestazione aziendale, fogli multipli,
# stile coerente con genera_report_direzione.py della dashboard PNRR gemella)
# --------------------------------------------------------------------------- #
RIGA_INIZIO_TABELLA = 6  # righe 1-4 intestazione aziendale, riga 5 vuota, riga 6 header tabella


def _intestazione_foglio(ws, titolo: str, sorgente_key: str, n_righe_dati: int) -> None:
    """Scrive l'intestazione aziendale nelle prime righe del foglio (la tabella dati
    è già stata scritta a partire da RIGA_INIZIO_TABELLA)."""
    ws["A1"] = config.AZIENDA_NOME
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = config.AZIENDA_UOC
    ws["A2"].font = Font(size=10, italic=True, color="666666")
    ws["A3"] = titolo
    ws["A3"].font = Font(bold=True, size=11)
    generato = pd.Timestamp.now().strftime("%d/%m/%Y %H:%M")
    sorgente = ETICHETTE_SORGENTE.get(sorgente_key, sorgente_key)
    ws["A4"] = f"Generato il {generato} · Sorgente dati: {sorgente} · {n_righe_dati} record"
    ws["A4"].font = Font(size=9, color="666666")


def _stila_header_tabella(ws, riga: int, n_col: int) -> None:
    for cella in ws[riga][:n_col]:
        cella.font = Font(bold=True, size=10, color="FFFFFF")
        cella.fill = PatternFill(start_color=COLORE_HEADER, end_color=COLORE_HEADER, fill_type="solid")
        cella.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws, riga_header: int, n_col: int, max_larghezza: int = 45) -> None:
    for idx in range(1, n_col + 1):
        lettera = get_column_letter(idx)
        larghezza = max(
            (len(str(ws.cell(row=r, column=idx).value or "")) for r in range(riga_header, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[lettera].width = min(max(larghezza + 2, 10), max_larghezza)


def _scrivi_foglio(writer, nome_foglio: str, titolo: str, df_dati: pd.DataFrame,
                   sorgente_key: str, colora_stato_vetusta: bool = False) -> None:
    riga_header = RIGA_INIZIO_TABELLA
    df_dati.to_excel(writer, sheet_name=nome_foglio, index=False, startrow=riga_header - 1)
    ws = writer.sheets[nome_foglio]
    _intestazione_foglio(ws, titolo, sorgente_key, len(df_dati))
    n_col = len(df_dati.columns)
    _stila_header_tabella(ws, riga_header, n_col)
    if colora_stato_vetusta and "stato_vetusta" in df_dati.columns:
        col_stato = df_dati.columns.get_loc("stato_vetusta") + 1
        for i, stato in enumerate(df_dati["stato_vetusta"], start=riga_header + 1):
            colore = COLORI_SEMAFORO_XLSX.get(stato)
            if colore:
                ws.cell(row=i, column=col_stato).fill = PatternFill(
                    start_color=colore, end_color=colore, fill_type="solid")
    _autosize(ws, riga_header, n_col)


def costruisci_report_direzione(d: pd.DataFrame, sorgente_key: str) -> bytes:
    """Report Excel multi-foglio per la Direzione: Executive Summary, Vetustà &
    Alert (con semaforo colorato), Elenco apparecchiature. Intestazione aziendale
    su ogni foglio. Stile coerente col report Direzione della dashboard PNRR."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # -- Executive Summary --------------------------------------------- #
        valore_tot = pd.to_numeric(d["valore_economico"], errors="coerce").sum(min_count=1)
        eta_media = d["anzianita_anni"].mean()
        riepilogo = pd.DataFrame({
            "Indicatore": [
                "Apparecchiature (selezione)", "Tipologie distinte", "Zone", "Sedi / UO",
                "Età media (anni)", "", "🔴 Vetuste (oltre soglia)",
                "🟡 In avvicinamento alla soglia", "🟢 OK", "", "Valore economico totale",
            ],
            "Valore": [
                len(d), d["tipologia"].nunique(), d["zona"].nunique(), d["sede"].nunique(),
                f"{eta_media:.1f}" if pd.notna(eta_media) else "n.d.", "",
                int((d["stato_vetusta"] == "Vetusto").sum()),
                int((d["stato_vetusta"] == "In avvicinamento").sum()),
                int((d["stato_vetusta"] == "OK").sum()), "",
                f"€ {valore_tot:,.0f}" if pd.notna(valore_tot) else "n.d. (non presente nel NSIS)",
            ],
        })
        _scrivi_foglio(writer, "Executive Summary", "Report Grandi Apparecchiature — Sintesi",
                       riepilogo, sorgente_key)

        # -- Vetustà & Alert (ordinata per urgenza, semaforo colorato) ------ #
        ordine = {"Vetusto": 0, "In avvicinamento": 1, "OK": 2, "n.d.": 3}
        vetusta = (d.assign(_o=d["stato_vetusta"].map(ordine))
                    .sort_values(["_o", "anzianita_anni"], ascending=[True, False])
                    [COLONNE_ELENCO_EXPORT])
        _scrivi_foglio(writer, "Vetustà e Alert", "Vetustà, contratti e scadenze — ordinato per urgenza",
                       vetusta, sorgente_key, colora_stato_vetusta=True)

        # -- Elenco completo -------------------------------------------------#
        elenco = d[COLONNE_ELENCO_EXPORT].sort_values(["zona", "sede", "tipologia"])
        _scrivi_foglio(writer, "Elenco Apparecchiature", "Elenco completo apparecchiature (selezione filtrata)",
                       elenco, sorgente_key)

    return buf.getvalue()


st.sidebar.title("🩻 Grandi Apparecchiature")
st.sidebar.caption("AUSL Toscana Nord Ovest — UOC Tecnologie")

sorgente = st.sidebar.radio(
    "Sorgente dati", ["NSIS (offline)", "API (cache)", "API live"], index=0,
    help="NSIS: sviluppo offline. API (cache): legge parco.parquet generato da "
         "'python sync.py' (veloce). API live: interroga ELM in tempo reale (lento "
         "sul parco completo; usare il limite).",
)
nsis_path = st.sidebar.text_input("Percorso file NSIS", config.NSIS_PATH_DEFAULT)

usa_api = sorgente == "API live"
usa_cache = sorgente == "API (cache)"
max_records = None
if usa_api:
    if st.sidebar.checkbox("Limita n. cespiti scaricati (consigliato in test)", value=True):
        max_records = st.sidebar.number_input(
            "Max cespiti", min_value=50, max_value=60000, value=500, step=50,
            help="Il parco completo conta decine di migliaia di record; con il rate "
                 "limit di 2 req/s un caricamento pieno richiede minuti.")

sorgente_key = "api" if usa_api else "cache" if usa_cache else "nsis"
try:
    df = carica(sorgente_key, nsis_path, max_records=max_records)
except Exception as exc:                                  # noqa: BLE001
    st.error(f"Impossibile caricare i dati dalla sorgente selezionata.\n\n{exc}")
    st.stop()

st.sidebar.divider()
if st.sidebar.button("🔄 Aggiorna dati", use_container_width=True):
    st.cache_data.clear()
    st.rerun()
st.sidebar.caption("Ricarica la sorgente corrente ignorando la cache.")

st.sidebar.divider()
st.sidebar.subheader("📊 Info dataset")
st.sidebar.metric("Record caricati", len(df))
st.sidebar.metric("Sedi / UO", df["sede"].nunique())
st.sidebar.metric("Zone", df["zona"].nunique())

# --------------------------------------------------------------------------- #
# Filtri
# --------------------------------------------------------------------------- #
st.sidebar.divider()
st.sidebar.header("Filtri")


def _ms(label, col):
    valori = sorted(v for v in df[col].dropna().unique())
    return st.sidebar.multiselect(label, valori, default=[])


f_azienda = _ms("Sotto-azienda", "azienda")
f_zona = _ms("Zona (distretto)", "zona")
f_tipo = _ms("Tipologia", "tipologia")
f_fabbr = _ms("Fabbricante", "fabbricante")
f_mod = _ms("Modalità acquisizione", "modalita_acquisizione")
solo_vetuste = st.sidebar.checkbox("Solo vetuste (oltre soglia)")
nascondi_accessori = st.sidebar.checkbox(
    "Nascondi consolle / accessori / iniettori",
    help="Consolle di comando, tavoli e iniettori sono censiti in ELM come cespiti "
         "a sé stanti, distinti dall'apparecchiatura principale (TAC/RM/mammografo/"
         "ecc.). Di default restano visibili come righe separate: spunta per "
         "mostrare solo le unità principali. Classificazione basata sul legame "
         "padre-figlio del cespite quando disponibile, altrimenti sul nome classe.")

d = df.copy()
if f_azienda:
    d = d[d["azienda"].isin(f_azienda)]
if f_zona:
    d = d[d["zona"].isin(f_zona)]
if f_tipo:
    d = d[d["tipologia"].isin(f_tipo)]
if f_fabbr:
    d = d[d["fabbricante"].isin(f_fabbr)]
if f_mod:
    d = d[d["modalita_acquisizione"].isin(f_mod)]
if solo_vetuste:
    d = d[d["vetusto"]]
if nascondi_accessori:
    d = d[d["componente"] != "Accessorio / Consolle"]

# --------------------------------------------------------------------------- #
# Intestazione + KPI
# --------------------------------------------------------------------------- #
st.title("Cruscotto Grandi Apparecchiature")
if sorgente_key == "nsis":
    st.caption("Sorgente: estrazione NSIS (offline). Valore economico e scadenze "
               "non presenti in questa sorgente — disponibili con l'API AT2.0.")

k1, k2, k3, k4, k5, k6, k7 = st.columns(7)
k1.metric("Apparecchiature", len(d))
k2.metric("Tipologie", d["tipologia"].nunique())
k3.metric("Sotto-aziende", d["azienda"].nunique())
k4.metric("Zone", d["zona"].nunique())
k5.metric("Sedi / UO", d["sede"].nunique())
eta_media = d["anzianita_anni"].mean()
k6.metric("Età media (anni)", f"{eta_media:.1f}" if pd.notna(eta_media) else "n.d.")
k7.metric("Vetuste (oltre soglia)", int(d["vetusto"].sum()))

valore_tot = pd.to_numeric(d["valore_economico"], errors="coerce").sum(min_count=1)
if pd.notna(valore_tot):
    st.metric("Valore economico totale (filtrato)", f"€ {valore_tot:,.0f}")

n_vet = int((d["stato_vetusta"] == "Vetusto").sum())
n_avv = int((d["stato_vetusta"] == "In avvicinamento").sum())
if n_vet:
    st.error(f"🔴 {n_vet} apparecchiature vetuste (oltre soglia)" +
             (f" · 🟡 {n_avv} in avvicinamento" if n_avv else ""))
elif n_avv:
    st.warning(f"🟡 {n_avv} apparecchiature in avvicinamento alla soglia di vetustà")

st.divider()

# --------------------------------------------------------------------------- #
# Sezioni
# --------------------------------------------------------------------------- #
tab_pan, tab_elenco, tab_vetusta, tab_export = st.tabs(
    ["Panoramica", "Elenco & dettaglio", "Vetustà & scadenze", "Export"])

with tab_pan:
    c1, c2 = st.columns(2)
    with c1:
        g = d["tipologia"].value_counts().reset_index()
        g.columns = ["tipologia", "n"]
        st.plotly_chart(px.bar(g, x="tipologia", y="n", title="Per tipologia",
                               text_auto=True), use_container_width=True)
    with c2:
        g = d["azienda"].value_counts().reset_index()
        g.columns = ["azienda", "n"]
        st.plotly_chart(px.bar(g, x="azienda", y="n", title="Per sotto-azienda",
                               text_auto=True), use_container_width=True)
    c3, c4 = st.columns(2)
    with c3:
        g = d["zona"].value_counts().reset_index()
        g.columns = ["zona", "n"]
        st.plotly_chart(px.bar(g, x="zona", y="n", title="Per zona (distretto)",
                               text_auto=True), use_container_width=True)
    with c4:
        g = d["fabbricante"].value_counts().reset_index()
        g.columns = ["fabbricante", "n"]
        st.plotly_chart(px.bar(g, x="fabbricante", y="n", title="Per fabbricante",
                               text_auto=True), use_container_width=True)
    c5, c6 = st.columns(2)
    with c5:
        dd = d.dropna(subset=["anzianita_anni"])
        st.plotly_chart(px.histogram(dd, x="anzianita_anni", nbins=20,
                                     title="Distribuzione anzianità (anni)"),
                        use_container_width=True)
    with c6:
        ac = d.dropna(subset=["anno_collaudo"])
        if not ac.empty:
            g = ac.groupby(ac["anno_collaudo"].astype(int)).size().reset_index(name="n")
            g.columns = ["anno_collaudo", "n"]
            st.plotly_chart(px.bar(g, x="anno_collaudo", y="n",
                                   title="Parco per anno di collaudo", text_auto=True),
                            use_container_width=True)

    st.divider()
    st.subheader("Distribuzione tipologie per sotto-azienda")
    st.caption("Dove sono ubicate le apparecchiature di ciascun tipo (es. le RM) tra "
               "le sotto-aziende. Riflette tutti i filtri attivi in sidebar.")

    tipologie_pan = sorted(d["tipologia"].dropna().unique())
    tipo_evid = st.selectbox(
        "Evidenzia tipologia", ["(nessuna)"] + tipologie_pan,
        help="Seleziona una tipologia (es. RM) per isolarne la distribuzione per "
             "sotto-azienda e zona.",
    )

    pivot_tipo_azienda = d.pivot_table(
        index="tipologia", columns="azienda", aggfunc="size", fill_value=0)
    fig_heat = px.imshow(
        pivot_tipo_azienda, text_auto=True, color_continuous_scale="Blues",
        aspect="auto", labels=dict(x="Sotto-azienda", y="Tipologia", color="N. apparecchiature"),
    )
    fig_heat.update_layout(height=max(300, 28 * len(pivot_tipo_azienda)))
    st.plotly_chart(fig_heat, use_container_width=True)

    if tipo_evid != "(nessuna)":
        de = d[d["tipologia"] == tipo_evid]
        ce1, ce2 = st.columns(2)
        with ce1:
            g = de["azienda"].value_counts().reset_index()
            g.columns = ["azienda", "n"]
            st.plotly_chart(px.bar(g, x="azienda", y="n",
                                   title=f"{tipo_evid} — per sotto-azienda", text_auto=True),
                            use_container_width=True)
        with ce2:
            g = de["zona"].value_counts().reset_index()
            g.columns = ["zona", "n"]
            st.plotly_chart(px.bar(g, x="zona", y="n",
                                   title=f"{tipo_evid} — per zona (distretto)", text_auto=True),
                            use_container_width=True)

with tab_elenco:
    st.subheader("Elenco apparecchiature")
    st.dataframe(
        d[["n_inventario", "tipologia", "componente", "fabbricante", "modello", "zona", "sede",
           "anno_collaudo", "anzianita_anni", "vetusto", "modalita_acquisizione",
           "valore_economico"]].style.format({"valore_economico": _euro}, na_rep="n.d."),
        use_container_width=True, hide_index=True,
    )
    st.subheader("Dettaglio per sede")
    sedi = sorted(s for s in d["sede"].dropna().unique())
    if sedi:
        sede_sel = st.selectbox("Sede / Unità Operativa", sedi)
        st.dataframe(d[d["sede"] == sede_sel], use_container_width=True, hide_index=True)

    st.divider()
    st.subheader("📍 Distribuzione geografica")
    st.info("La mappa delle sedi non è ancora disponibile: servono le coordinate GPS "
            "delle strutture (non presenti né nel NSIS né nell'attuale sorgente). "
            "TODO non bloccante — da aggiungere con un'anagrafica sedi georeferenziata.")

with tab_vetusta:
    st.subheader("Stato vetustà del parco (filtrato)")
    s1, s2, s3 = st.columns(3)
    s1.metric("🟢 OK", int((d["stato_vetusta"] == "OK").sum()))
    s2.metric("🟡 In avvicinamento", int((d["stato_vetusta"] == "In avvicinamento").sum()))
    s3.metric("🔴 Vetuste", int((d["stato_vetusta"] == "Vetusto").sum()))

    ordine = {"Vetusto": 0, "In avvicinamento": 1, "OK": 2, "n.d.": 3}
    tabella = d.assign(_o=d["stato_vetusta"].map(ordine)).sort_values(
        ["_o", "anzianita_anni"], ascending=[True, False])
    vista = tabella[["stato_vetusta", "n_inventario", "tipologia", "fabbricante",
                     "modello", "zona", "sede", "anno_collaudo", "anzianita_anni",
                     "soglia_vetusta", "valore_economico"]]

    colori = {"Vetusto": "background-color:#f8d7da",
              "In avvicinamento": "background-color:#fff3cd",
              "OK": "background-color:#d1e7dd"}
    styled = (vista.style
              .map(lambda v: colori.get(v, ""), subset=["stato_vetusta"])
              .format({"valore_economico": _euro}, na_rep="n.d."))
    st.dataframe(styled, use_container_width=True, hide_index=True)

    scad_cols = ["scadenza_contratto", "scadenza_vse", "scadenza_cq"]
    if d[scad_cols].notna().any().any():
        st.subheader("Scadenze contratti / verifiche")
        st.dataframe(d[d[scad_cols].notna().any(axis=1)],
                     use_container_width=True, hide_index=True)
    else:
        st.info("Scadenze contratti/verifiche non disponibili in questa sorgente "
                "(saranno popolate dall'API AT2.0).")

with tab_export:
    st.subheader("📊 Report Direzione")
    st.caption("Excel multi-foglio con intestazione aziendale: Executive Summary, "
               "Vetustà & Alert (con semaforo colorato) ed Elenco apparecchiature. "
               "Riflette i filtri correnti in sidebar.")
    report_bytes = costruisci_report_direzione(d, sorgente_key)
    st.download_button(
        "⬇️ Scarica Report Direzione (Excel)", data=report_bytes,
        file_name=f"report_grandi_apparecchiature_{pd.Timestamp.today():%Y%m%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.divider()
    st.subheader("Esportazione rapida (dati grezzi)")
    st.caption("Foglio unico con tutte le colonne, senza formattazione — per elaborazioni successive.")
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        d.to_excel(writer, index=False, sheet_name="Grandi Apparecchiature")
    st.download_button(
        "⬇️ Scarica Excel grezzo (dati filtrati)", data=buf.getvalue(),
        file_name="grandi_apparecchiature.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
